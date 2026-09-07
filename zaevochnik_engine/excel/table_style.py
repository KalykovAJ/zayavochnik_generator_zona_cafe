from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, Protection
from openpyxl.formatting.rule import FormulaRule

from zaevochnik_engine.excel.data_validation import build_qty_validations
from zaevochnik_engine.excel.layout import autofit_column_width


def _build_style_kit(excel_styles: dict) -> dict:
    """Собирает переиспользуемые объекты стилей (Font/Fill/Border/Alignment)
    из EXCEL_STYLES конкретной сети."""
    colors = excel_styles["colors"]
    fonts_cfg = excel_styles["fonts"]
    align_cfg = excel_styles["alignments"]

    return {
        "header_fill": PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid"),
        "header_font": Font(**fonts_cfg["header"]),

        "pack_fill": PatternFill(start_color=colors["bg_pack"], fill_type="solid"),
        "default_fill": PatternFill(fill_type=None),

        "regular_font": Font(**fonts_cfg["regular"]),
        "bold_font": Font(**fonts_cfg["bold"]),

        "align_center": Alignment(**align_cfg["default"]),
        "align_left": Alignment(**align_cfg["text_left"]),
        "align_header": Alignment(**align_cfg["header"]),

        "border_data": Border(
            left=Side(border_style="thin", color=colors["border"]),
            right=Side(border_style="thin", color=colors["border"]),
            top=Side(border_style="thin", color=colors["border"]),
            bottom=Side(border_style="thin", color=colors["border"])
        ),
        "border_header": Border(
            left=Side(border_style="thin", color=colors["border"]),
            right=Side(border_style="thin", color=colors["border"]),
            top=Side(border_style="thin", color=colors["border"]),
            bottom=Side(border_style="medium", color=colors["primary"])
        )
    }


def _apply_header_row_style(worksheet, start_row: int, styles: dict, row_height: int):
    """Оформляет строку-шапку таблицы (строка header_row, напр. строка 5)."""
    worksheet.row_dimensions[start_row].height = row_height
    for col_idx in range(1, worksheet.max_column + 1):
        cell = worksheet.cell(row=start_row, column=col_idx)
        cell.fill = styles["header_fill"]
        cell.font = styles["header_font"]
        cell.border = styles["border_header"]
        cell.alignment = styles["align_header"]


def _apply_data_rows_style(worksheet, start_row: int, end_row: int, column_mapping: dict, styles: dict,
                           row_height: int, dv_unit, dv_pack):
    """Оформляет строки с данными: заливка по типу заказа (упаковка/штуки),
    Data Validation, защита (редактировать можно только колонку количества),
    выравнивание."""
    headers = {str(worksheet.cell(row=start_row, column=i).value).strip(): i for i in
               range(1, worksheet.max_column + 1)}
    type_idx = headers.get(column_mapping["type_col"])
    qty_idx = headers.get(column_mapping["qty_col"])

    for row_idx in range(start_row + 1, end_row + 1):
        worksheet.row_dimensions[row_idx].height = row_height

        is_pack, is_unit = False, False

        if type_idx:
            cell_type = worksheet.cell(row=row_idx, column=type_idx)
            current_type_val = str(cell_type.value or "").strip()
            type_text = current_type_val.lower()

            is_pack = "упаковк" in type_text
            is_unit = "штук" in type_text

            if (is_pack or is_unit) and "➔" not in current_type_val:
                cell_type.value = f"{current_type_val} ➔"

        current_fill = styles["pack_fill"] if (is_pack or is_unit) else styles["default_fill"]

        if qty_idx:
            cell_qty = worksheet.cell(row=row_idx, column=qty_idx)

            if is_pack:
                dv_pack.add(cell_qty)
                cell_qty.value = None
                # Защита дефолтного стиля ячейки от Ctrl+V через условное форматирование
                style_rule = FormulaRule(formula=['1=1'], fill=styles["default_fill"], font=styles["regular_font"],
                                         border=styles["border_data"])
                worksheet.conditional_formatting.add(cell_qty.coordinate, style_rule)
            elif is_unit:
                dv_unit.add(cell_qty)
                cell_qty.value = None
                style_rule = FormulaRule(formula=['1=1'], fill=styles["default_fill"], font=styles["regular_font"],
                                         border=styles["border_data"])
                worksheet.conditional_formatting.add(cell_qty.coordinate, style_rule)

        for col_idx in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            header_name = str(worksheet.cell(row=start_row, column=col_idx).value or "").strip()
            header_val_lower = header_name.lower()

            if col_idx == qty_idx:
                cell.fill = styles["default_fill"]
            elif current_fill.fill_type:
                cell.fill = current_fill

            cell.font = styles["bold_font"] if "итого" in header_val_lower else styles["regular_font"]
            cell.border = styles["border_data"]

            # В обычном режиме редактировать можно только ячейку количества
            cell.protection = Protection(locked=(col_idx != qty_idx))

            # Раньше здесь сравнивалось "header_name == 'Наименование'" —
            # реальная колонка в справочнике называется "Наименование товара",
            # так что условие никогда не срабатывало и текст не выравнивался
            # по левому краю. startswith() устойчив к любому варианту названия.
            if header_name.startswith("Наименование"):
                cell.alignment = styles["align_left"]
            else:
                cell.alignment = styles["align_center"]


def apply_table_row_styles(worksheet, start_row: int, end_row: int, column_mapping: dict,
                           excel_styles: dict, validation_prompts: dict):
    """Точка входа: применяет стили оформления шапки таблицы и строк данных."""
    heights = excel_styles["row_heights"]
    styles = _build_style_kit(excel_styles)
    dv_unit, dv_pack = build_qty_validations(worksheet, validation_prompts)

    _apply_header_row_style(worksheet, start_row, styles, row_height=heights["header"])
    _apply_data_rows_style(
        worksheet, start_row, end_row, column_mapping, styles,
        row_height=heights["data"], dv_unit=dv_unit, dv_pack=dv_pack
    )

    autofit_column_width(
        worksheet,
        col_indices=range(1, worksheet.max_column + 1),
        start_row=1,
        end_row=end_row,
        min_width=15,
        padding=4
    )
