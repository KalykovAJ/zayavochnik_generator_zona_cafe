import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter

from zaevochnik_engine.excel.formulas import get_excel_col_letter
from zaevochnik_engine.excel.layout import autofit_column_width


def _apply_brand_block(worksheet, max_col: int, cfg: dict, start_row: int, end_row: int,
                       weight_column_name: str):
    """Заливка/границы/тексты брендированного блока (строки 1-4)."""
    bg_fill = PatternFill(start_color=cfg["bg_color"], end_color=cfg["bg_color"], fill_type="solid")

    font_company = Font(name=cfg["font_name"],
                        size=cfg.get("company_name_font_size", cfg["company_font_size"]),
                        bold=True, color=cfg["text_color_white"])
    font_yellow = Font(name=cfg["font_name"], size=cfg["label_font_size"], bold=True, color=cfg["text_color_yellow"])
    font_white = Font(name=cfg["font_name"], size=cfg["label_font_size"], bold=True, color=cfg["text_color_white"])

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    border_color = cfg.get("border", "000000")
    thin_side = Side(border_style="thin", color=border_color)

    for r in range(1, 5):
        for c in range(1, max_col + 1):
            cell = worksheet.cell(row=r, column=c)
            cell.fill = bg_fill
            cell.protection = Protection(locked=True)
            cell.border = Border(
                top=thin_side if r == 1 else None,
                bottom=thin_side if r == 4 else None,
                left=thin_side if c == 1 else None,
                right=thin_side if c == max_col else None
            )

    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = f"Последнее обновление заявочника: {datetime.datetime.now().strftime('%d.%m.%Y')}"
    worksheet["A1"].font = font_yellow
    worksheet["A1"].alignment = align_left

    # Правая часть шапки (подпись + сумма веса) считается динамически, чтобы
    # оставаться на месте независимо от текущего числа колонок таблицы.
    col_lbl_letter = get_column_letter(max_col - 1)
    col_val_letter = get_column_letter(max_col)

    # Блок названия компании: правая граница считается динамически - до "F"
    # включительно, но не дальше двух колонок перед блоком веса, чтобы не
    # пересекаться с ним на узких таблицах.
    company_end_col = max(min(6, max_col - 2), 3)
    company_end_letter = get_column_letter(company_end_col)

    worksheet.merge_cells(f"C1:{company_end_letter}4")
    worksheet["C1"] = cfg["company_name"]
    worksheet["C1"].font = font_company
    worksheet["C1"].alignment = align_center

    worksheet["A2"] = cfg["fuel_station"]
    worksheet["A2"].font = font_white
    worksheet["A2"].alignment = align_right

    worksheet["B2"] = "Введите номер"
    worksheet["B2"].font = font_white
    worksheet["B2"].alignment = align_left
    worksheet["B2"].protection = Protection(locked=False)

    worksheet["A3"] = "Дата заявки:"
    worksheet["A3"].font = font_white
    worksheet["A3"].alignment = align_right

    worksheet["B3"] = "Введите дату"
    worksheet["B3"].font = font_white
    worksheet["B3"].alignment = align_left
    worksheet["B3"].number_format = "DD.MM.YYYY"
    worksheet["B3"].protection = Protection(locked=False)

    worksheet.merge_cells(f"{col_lbl_letter}1:{col_lbl_letter}4")
    worksheet[f"{col_lbl_letter}1"] = "Общий вес:"
    worksheet[f"{col_lbl_letter}1"].font = font_white
    worksheet[f"{col_lbl_letter}1"].alignment = align_center

    worksheet.merge_cells(f"{col_val_letter}1:{col_val_letter}4")
    f_total_weight_letter = get_excel_col_letter(worksheet, start_row, weight_column_name)
    worksheet[f"{col_val_letter}1"] = f"=SUM({f_total_weight_letter}{start_row + 1}:{f_total_weight_letter}{end_row})"
    worksheet[f"{col_val_letter}1"].font = font_white
    worksheet[f"{col_val_letter}1"].alignment = align_center
    worksheet[f"{col_val_letter}1"].number_format = '#,##0.00" кг"'


def _apply_row_protection(worksheet, start_row: int, end_row: int, column_mapping: dict, max_col: int):
    """Защита строк данных: редактировать можно только колонку количества.

    Раньше здесь ещё проверялись правила полной блокировки строки по
    статусу товара/типу поставки - в справочнике "Зона кафе" такой
    колонки ("Статус") нет вовсе, и значения типа заказа, на которые
    реагировали эти правила ("напрямую", "приостановлена" и т.п.),
    в нём никогда не встречаются, так что эта ветка была мёртвым кодом
    и удалена вместе с excel/lock_rules.py.
    """
    qty_letter = get_excel_col_letter(worksheet, start_row, column_mapping["qty_col"])

    for row in range(start_row + 1, end_row + 1):
        worksheet[f"{qty_letter}{row}"].protection = Protection(locked=False)
        for col_idx in range(1, max_col + 1):
            if get_column_letter(col_idx) != qty_letter:
                worksheet.cell(row=row, column=col_idx).protection = Protection(locked=True)


def _autofit_header_layout(worksheet, max_col: int):
    """Автоподбор ширины колонок A/B и высоты строк 1-4 брендированной шапки."""
    autofit_column_width(worksheet, col_indices=[1, 2], start_row=1, end_row=4,
                         min_width=10, padding=5, max_width=20)

    for r in range(2, 5):
        max_row_height = 22
        for c in range(1, max_col + 1):
            cell = worksheet.cell(row=r, column=c)
            if cell.value is not None and not str(cell.value).startswith("=") and c <= 2:
                val_str = str(cell.value)
                col_letter = get_column_letter(c)
                col_width = worksheet.column_dimensions[col_letter].width or 12
                if len(val_str) > col_width:
                    lines_count = (len(val_str) // int(col_width)) + 1
                    calculated_height = lines_count * 16
                    max_row_height = max(max_row_height, calculated_height)
        worksheet.row_dimensions[r].height = max_row_height

    worksheet.row_dimensions[1].height = 24


def apply_top_header_and_protection(worksheet, start_row: int, end_row: int, column_mapping: dict,
                                    weight_column_name: str, excel_styles: dict):
    """
    Формирует и стилизует верхнюю шапку (строки 1-4), проставляет защиту
    строк данных (редактируется только колонка количества) и включает
    защиту листа.
    """
    cfg = excel_styles["top_header"]
    max_col = worksheet.max_column

    _apply_brand_block(worksheet, max_col, cfg, start_row, end_row, weight_column_name)
    _apply_row_protection(worksheet, start_row, end_row, column_mapping, max_col)
    _autofit_header_layout(worksheet, max_col)

    worksheet.protection.password = cfg["sheet_password"]
    worksheet.protection.selectLockedCells = True
    worksheet.protection.selectUnlockedCells = False
    worksheet.protection.enable()
