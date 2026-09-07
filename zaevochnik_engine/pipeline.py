from typing import List, Optional, Union

import pandas as pd
from openpyxl.utils import get_column_letter

from zaevochnik_engine.data.reader import read_source
from zaevochnik_engine.data.filterer import filter_by_network
from zaevochnik_engine.data.column_selector import select_final_columns, add_generated_columns
from zaevochnik_engine.excel.writer import write_dataframe
from zaevochnik_engine.excel.formulas import apply_dynamic_formulas
from zaevochnik_engine.excel.table_style import apply_table_row_styles
from zaevochnik_engine.excel.top_header import apply_top_header_and_protection


def build_zaevochnik(
        source_path: str,
        output_path: str,
        source_sheet: Union[str, List[str]],
        sheet_name: str,
        header_row: int,
        filter_col: Optional[Union[str, List[str]]],
        filter_val,
        columns_to_keep: List[str],
        column_mapping: dict,
        weight_column_name: str,
        excel_styles: dict,
        validation_prompts: dict,
):
    """Изолированный бизнес-процесс сборки заявочника.

    source_sheet - название листа справочника-источника, из которого строится
    ЭТОТ заявочник (например, "Заморозка" или "Кофе"). Какие заявочники
    (сочетания "категория товара" x "сеть") вообще нужно собрать, задаётся не
    здесь и не в коде, а в листе "Задания" файла настроек - см.
    zaevochnik_engine/config_loader.py и run_zaevochnik.py.

    columns_to_keep - БЕЛЫЙ список колонок исходника, которые должны попасть
    в финальный заявочник (см. data/column_selector.py). Флаговые колонки
    сетей (МП/БП/ПН/...) сюда не включаются и поэтому никогда не просачиваются
    в заявочник, даже если в справочник добавится ещё одна сеть.
    """
    print(f"Шаг 1: Чтение листа(ов) '{source_sheet}' и фильтрация по сети...")
    df = read_source(source_path, sheet_names=source_sheet)
    df = filter_by_network(df, filter_col, filter_val)

    print("Шаг 2: Отбор колонок (белый список) и добавление служебных полей...")
    df = select_final_columns(df, columns_to_keep)
    df = add_generated_columns(
        df,
        qty_col=column_mapping["qty_col"],
        total_col=column_mapping["total_col"],
        weight_total_col=weight_column_name
    )

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)
        worksheet = writer.sheets[sheet_name]

        print("Шаг 3: Запись данных в Excel...")
        total_rows = write_dataframe(worksheet, df, start_row=header_row)
        start_data_row = header_row + 1

        print("Шаг 4: Просчет и запись формул на структуру таблицы...")
        apply_dynamic_formulas(
            worksheet=worksheet,
            start_row=start_data_row,
            end_row=total_rows,
            header_row=header_row,
            column_mapping=column_mapping,
            weight_column_name=weight_column_name
        )

        print("Шаг 5: Наложение стилей, подсказок и валидации...")
        apply_table_row_styles(
            worksheet=worksheet,
            start_row=header_row,
            end_row=total_rows,
            column_mapping=column_mapping,
            excel_styles=excel_styles,
            validation_prompts=validation_prompts
        )

        print("Шаг 6: Стилизация брендированной шапки макета и включение защиты листа...")
        apply_top_header_and_protection(
            worksheet=worksheet,
            start_row=header_row,
            end_row=total_rows,
            column_mapping=column_mapping,
            weight_column_name=weight_column_name,
            excel_styles=excel_styles
        )

        max_col_letter = get_column_letter(worksheet.max_column)
        worksheet.auto_filter.ref = f"A{header_row}:{max_col_letter}{total_rows}"
        worksheet.freeze_panes = f"A{start_data_row}"

    print(f"🎉 Процесс завершен! Заявочник обновлен и сохранен: {output_path}")
